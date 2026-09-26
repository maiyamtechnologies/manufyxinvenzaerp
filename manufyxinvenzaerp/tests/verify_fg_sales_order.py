"""Sales Order side of the finished-goods Kg / Nos plan (sep14 plan, package A2).

An FG line is ordered in Kg (its Quantity) and carries its piece count as Qty (Nos).
The upload sheet gives each drawing two customer weights -- one piece, and all of
them -- and Verify Raw Materials refuses to pass until:

  per drawing   Cust Weight (per Nos) x Total Qty = Cust Weight (Total)
  per FG item   the drawings' Totals add up to the line Kg, their Total Qtys to the
                line Nos, and no drawing names an FG item the order does not sell.

Create Drawing checks the pass on the server, and a Drawing List row that already
has its Drawing cannot have its weights or Nos changed except to follow the Drawing
(Update Customer Weight).

The whole run is ONE database transaction that is rolled back at the end:
frappe.db.commit is switched off while it runs, so the ZZFG- Sales Order, its staged
rows and the Drawing made from them never reach the database for anyone else, and no
naming series moves. The sheet is fed to the real parser from a temp file by standing
in for the attachment lookup, so nothing is written under the site's files either.

Run: bench --site manufact execute manufyxinvenzaerp.tests.verify_fg_sales_order.run
"""

import io
import os
import tempfile

import frappe
from frappe.utils import flt, nowdate

checks = []
TAG = frappe.generate_hash(length=6).upper()
FG = "Fabricated Structurs"
FOREIGN_FG = "FINGOODS001"


def check(label, got, want):
    ok = got == want
    checks.append(ok)
    print("  %-4s %-66s got=%r want=%r" % ("OK" if ok else "FAIL", label, got, want))


def _plain(text):
    return frappe.utils.strip_html(text or "").replace("&middot;", "·")


def _throws(fn):
    try:
        fn()
        return None
    except frappe.ValidationError as e:
        return _plain(str(e)) or "thrown"


def _messages():
    out = []
    for m in (frappe.local.message_log or []):
        if isinstance(m, str):
            m = frappe.parse_json(m)
        out.append(_plain((m.get("title") or "") + " | " + (m.get("message") or "")))
    return out


def _write_sheet(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(list(r))
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return path


def run():
    real_commit = frappe.db.commit
    frappe.db.commit = lambda *a, **k: None
    try:
        _run()
    finally:
        frappe.db.rollback()
        frappe.db.commit = real_commit
        print("rolled back: nothing from this run is left in the database")

    print()
    print("=== SUMMARY ===")
    if all(checks):
        print("ALL %d CHECKS PASSED" % len(checks))
    else:
        print("%d of %d CHECKS FAILED" % (checks.count(False), len(checks)))


def _run():
    from manufyxinvenzaerp.drawing_management import so_drawing_import as imp
    from manufyxinvenzaerp.production_management.fg_stock import is_fg_item

    print("=== fixtures ===")
    customer = frappe.db.get_value("Customer", {"disabled": 0}, "name")
    company = frappe.db.get_value("Company", {}, "name")
    material = frappe.db.get_value(
        "Item", {"custom_parent_item_group": "Structurals", "custom_unit_weight": [">", 0],
                 "disabled": 0}, "name")
    print("   customer=%r company=%r material=%r" % (customer, company, material))
    check("the FG item is finished goods", is_fg_item(FG), True)
    check("the second FG item is finished goods too", is_fg_item(FOREIGN_FG), True)
    if not (customer and company and material):
        print("   missing fixtures -- stopped")
        return

    print()
    print("=== 1. template: both customer weights, sample 5 x 50 = 250 ===")
    import openpyxl
    frappe.response.clear()
    imp.download_bom_template()
    wb = openpyxl.load_workbook(io.BytesIO(frappe.response["filecontent"]), data_only=True)
    trows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    frappe.response.clear()
    header = list(trows[0])
    check("headers in plan order",
          header[3:9], ["FG Item", "Total Qty", "Cust Weight (per Nos)", "Cust Weight (Total)",
                        "Nature of Work", "Rate Schedule"])
    check("old weight header is gone", "Total Weight (KG)" in header, False)
    s = dict(zip(header, trows[1]))
    check("sample is 5 Nos x 50 = 250",
          (s["Total Qty"], s["Cust Weight (per Nos)"], s["Cust Weight (Total)"]), (5, 50, 250))

    print()
    print("=== 2. parser: new and old headers ===")
    cdn = "ZZFG-CDN-%s-1" % TAG
    duno = "ZZFG-%s-1" % TAG
    body = ["ZZFG", cdn, duno, FG, 5, 50, 250, "", "", "1", material, "", 0, 0, 1000, 1]
    path = _write_sheet([header, body])
    parsed = imp._parse_excel(path)[cdn]
    check("new headers: per Nos and Total read",
          (parsed["weight_per_pcs"], parsed["total_weight"]), (50.0, 250.0))
    old_header = [h for h in header]
    old_header[old_header.index("Cust Weight (per Nos)")] = "Weight per Pcs (KG)"
    old_header[old_header.index("Cust Weight (Total)")] = "Total Weight (KG)"
    old_path = _write_sheet([old_header, body])
    try:
        old = imp._parse_excel(old_path)[cdn]
        check("old headers: per Nos and Total read",
              (old["weight_per_pcs"], old["total_weight"]), (50.0, 250.0))
    finally:
        os.unlink(old_path)

    print()
    print("=== 3. load the sheet onto a 250 Kg / 5 Nos order ===")
    so = frappe.get_doc({
        "doctype": "Sales Order",
        "customer": customer,
        "company": company,
        "po_no": "ZZFG-PO-%s" % TAG,
        "transaction_date": nowdate(),
        "delivery_date": nowdate(),
        "items": [{"item_code": FG, "qty": 250, "uom": "Kg", "rate": 10,
                   "custom_sec_qty": 5, "delivery_date": nowdate()}],
    })
    so.flags.name_set = True
    so.name = "ZZFG-SO-%s" % TAG
    so.insert(ignore_permissions=True)
    frappe.db.set_value("Sales Order", so.name, "custom_bom_excel_file",
                        "/private/files/zzfg-%s.xlsx" % TAG, update_modified=False)
    real_get_path = imp._get_file_path
    imp._get_file_path = lambda url: path
    try:
        res = imp.parse_bom_excel(so.name)
    finally:
        imp._get_file_path = real_get_path
        os.unlink(path)
    check("one drawing staged", res["drawing_count"], 1)
    staged = frappe.db.get_value("Sales Order DUNO Item", {"parent": so.name},
                                 ["name", "weight_per_pcs", "total_weight", "total_quantity"],
                                 as_dict=True)
    check("staged per Nos / Total / Nos",
          (flt(staged.weight_per_pcs, 3), flt(staged.total_weight, 3), flt(staged.total_quantity, 3)),
          (50.0, 250.0, 5.0))
    line = frappe.db.get_value("Sales Order Item", {"parent": so.name}, "name")

    def verify():
        frappe.local.message_log = []
        r = imp.verify_raw_materials(so.name)
        return r["verified"], [_plain(i) for i in r["issues"]]

    def flag():
        return frappe.db.get_value("Sales Order", so.name, "custom_raw_materials_verified")

    print()
    print("=== 4. Verify Raw Materials ===")
    ok, issues = verify()
    for i in issues:
        print("      •", i[:150])
    check("5 x 50 = 250 on a 250 Kg / 5 Nos line verifies", ok, True)
    check("...and raises the flag", flag(), 1)

    frappe.db.set_value("Sales Order DUNO Item", staged.name, "total_weight", 245)
    ok, issues = verify()
    row_issue = [i for i in issues if i.startswith("Drawing List row 1 ·") and "× 5 Nos = 250" in i]
    check("a Total of 245 fails", ok, False)
    check("...naming the row, per Nos x Nos and the Total", len(row_issue), 1)
    if row_issue:
        print("       ->", row_issue[0][:160])
    check("...and the flag drops", flag(), 0)
    frappe.db.set_value("Sales Order DUNO Item", staged.name, "total_weight", 250)

    frappe.db.set_value("Sales Order DUNO Item", staged.name, "weight_per_pcs", 0)
    ok, issues = verify()
    check("a missing per Nos fails",
          ok is False and any("Cust Weight (per Nos) is missing" in i for i in issues), True)
    frappe.db.set_value("Sales Order DUNO Item", staged.name, "weight_per_pcs", 50)

    frappe.db.set_value("Sales Order Item", line, "qty", 240)
    ok, issues = verify()
    kg_issue = [i for i in issues if i.startswith("Items row 1 ·")]
    check("a 240 Kg line fails", ok, False)
    check("...ordered vs drawings in Kg and Nos, with the difference",
          bool(kg_issue) and all(t in kg_issue[0] for t in
                                 ("planned 250 Kg / 5 Nos", "Items table has 240 Kg / 5 Nos", "Drawing List has 10 Kg more")), True)
    if kg_issue:
        print("       ->", kg_issue[0][:170])
    frappe.db.set_value("Sales Order Item", line, "qty", 250)

    # Above the drawings is a warning only (2026-09-26): the customer's weight may sit
    # over the planned one. Measured as "no Items-row issue", not as a pass, so it
    # holds whatever else this fixture's sheet trips over.
    frappe.db.set_value("Sales Order Item", line, "qty", 260)
    r = imp.verify_raw_materials(so.name)
    check("a 260 Kg line (above planned) does not block",
          any(_plain(i).startswith("Items row 1") for i in r["issues"]), False)
    check("...it is a warning naming the excess",
          any("Items table has 10 Kg more" in _plain(w) for w in r["warnings"]), True)
    frappe.db.set_value("Sales Order Item", line, "qty", 250)

    frappe.db.set_value("Sales Order Item", line, "custom_sec_qty", 4)
    ok, issues = verify()
    check("a 4 Nos line fails", ok, False)
    check("...naming the Nos difference",
          any("250 Kg / 4 Nos" in i and "Drawing List has 1 Nos more" in i and "Kg more" not in i for i in issues), True)
    frappe.db.set_value("Sales Order Item", line, "custom_sec_qty", 5)

    frappe.db.set_value("Sales Order DUNO Item", staged.name, "item", FOREIGN_FG)
    ok, issues = verify()
    check("a drawing for an FG item the order does not sell fails", ok, False)
    check("...saying it is not on the items table",
          any(FOREIGN_FG in i and "not on this Sales Order's items table" in i for i in issues), True)
    frappe.db.set_value("Sales Order DUNO Item", staged.name, "item", FG)

    ok, issues = verify()
    check("put back, it verifies again", (ok, issues), (True, []))

    print()
    print("=== 5. Sales Order validate keeps the flag honest ===")
    doc = frappe.get_doc("Sales Order", so.name)
    doc.items[0].qty = 251
    frappe.local.message_log = []
    doc.save(ignore_permissions=True)
    msgs = _messages()
    check("changing the FG line Kg clears the flag", flag(), 0)
    check("...and says why", any("Verification Cleared" in m for m in msgs), True)
    check("an orange warning names the mismatch on save",
          any("Planned weight differs" in m and "251 Kg / 5 Nos" in m and "Items table has 1 Kg more" in m for m in msgs), True)
    doc = frappe.get_doc("Sales Order", so.name)
    doc.items[0].qty = 250
    frappe.local.message_log = []
    doc.save(ignore_permissions=True)
    check("matching again: no mismatch warning",
          any("Planned weight differs" in m for m in _messages()), False)

    verify()
    doc = frappe.get_doc("Sales Order", so.name)
    doc.items[0].custom_sec_qty = 6
    doc.save(ignore_permissions=True)
    check("changing the FG line Nos clears the flag", flag(), 0)
    doc = frappe.get_doc("Sales Order", so.name)
    doc.items[0].custom_sec_qty = 5
    doc.save(ignore_permissions=True)

    verify()
    doc = frappe.get_doc("Sales Order", so.name)
    doc.custom_duno_items[0].weight_per_pcs = 51
    doc.save(ignore_permissions=True)
    check("changing a pending Drawing List weight clears the flag", flag(), 0)
    doc = frappe.get_doc("Sales Order", so.name)
    doc.custom_duno_items[0].weight_per_pcs = 50
    doc.save(ignore_permissions=True)

    verify()
    doc = frappe.get_doc("Sales Order", so.name)
    doc.remarks = "ZZFG unrelated edit"
    doc.save(ignore_permissions=True)
    check("an unrelated edit keeps the flag", flag(), 1)

    frappe.db.set_value("Sales Order", so.name, "custom_raw_materials_verified", 0)
    doc = frappe.get_doc("Sales Order", so.name)
    doc.custom_raw_materials_verified = 1
    doc.save(ignore_permissions=True)
    check("a flag raised by a save, not by Verify, is put back", flag(), 0)

    print()
    print("=== 6. Create Drawing checks the pass on the server ===")
    err = _throws(lambda: imp.create_drawings_from_import(so.name, 0, 30))
    check("calling Create Drawing directly while unverified is refused",
          bool(err) and "not verified" in err, True)
    check("...and nothing was created",
          frappe.db.count("Drawing", {"sales_order": so.name}), 0)

    verify()
    res = imp.create_drawings_from_import(so.name, 0, 30)
    made = [r for r in res["results"] if r["status"] == "success"]
    for r in res["results"]:
        if r["status"] != "success":
            print("      failed:", _plain(r.get("error"))[:160])
    check("verified, the drawing is created", len(made), 1)
    if not made:
        return
    drawing = frappe.db.get_value(
        "Drawing", made[0]["drawing"],
        ["customer_provided_wt", "weight_per_pcs", "no_of_qty_to_manufacture"], as_dict=True)
    check("Drawing Cust Weight (Total) = 250 and per Nos = 50, 5 Nos",
          (flt(drawing.customer_provided_wt, 3), flt(drawing.weight_per_pcs, 3),
           flt(drawing.no_of_qty_to_manufacture, 3)), (250.0, 50.0, 5.0))

    print()
    print("=== 7. a row with a Drawing is locked ===")
    for field, value in (("total_weight", 260), ("weight_per_pcs", 52), ("total_quantity", 6)):
        doc = frappe.get_doc("Sales Order", so.name)
        setattr(doc.custom_duno_items[0], field, value)
        err = _throws(lambda: doc.save(ignore_permissions=True))
        check("editing %s on the locked row is refused" % field,
              bool(err) and "Update Customer Weight" in err, True)

    doc = frappe.get_doc("Sales Order", so.name)
    doc.remarks = "ZZFG untouched row"
    check("saving with the row untouched is fine",
          _throws(lambda: doc.save(ignore_permissions=True)), None)

    # Update Customer Weight writes the Drawing first and then the same figures to
    # the row. Stood in for directly: the row may follow its Drawing.
    frappe.db.set_value("Drawing", made[0]["drawing"],
                        {"weight_per_pcs": 60, "customer_provided_wt": 300},
                        update_modified=False)
    doc = frappe.get_doc("Sales Order", so.name)
    doc.custom_duno_items[0].weight_per_pcs = 60
    doc.custom_duno_items[0].total_weight = 300
    frappe.local.message_log = []
    check("the row may follow its Drawing (Update Customer Weight)",
          _throws(lambda: doc.save(ignore_permissions=True)), None)
    check("...and the new line difference shows in orange",
          any("Planned weight differs" in m and "Drawing List has 50 Kg more" in m for m in _messages()), True)
